"""Build the 2 macro-kit IC slides on top of the 'French and Italian spreads briefing' template."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]

# Template palette
BG, CARD, BORDER = "F5F2EC", "FBF9F5", "DDD6CA"
NAVY, BLUE, RED = "17202B", "2B5C8A", "A8322B"
GREY, MUTED, ONDARK, ONDARK_MUTED = "5F6670", "6B7078", "F5F2EC", "B9BEC6"
SERIF, SANS = "Source Serif 4", "Inter"

prs = Presentation(SRC)

# Drop every template slide; keep master, layout and theme
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

layout = prs.slide_layouts[0]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide():
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    return s


def box(s, x, y, w, h, fill=CARD, line=BORDER, radius=0.04):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rule(s, x, y, w, color=BORDER):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(0.5)
    return ln


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of paragraphs; each paragraph is a list of runs (str, opts) or a single (str, opts)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        if isinstance(para, tuple):
            para = [para]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        first = para[0][1]
        p.line_spacing = Pt(first.get("lh", first.get("sz", 12) * 1.4))
        if i > 0:
            p.space_before = Pt(first.get("before", 0))
        for t, o in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("sz", 12))
            f.bold = o.get("b", False)
            f.italic = o.get("i", False)
            f.name = o.get("font", SANS)
            f.color.rgb = rgb(o.get("c", NAVY))
            if o.get("spc"):
                r._r.get_or_add_rPr().set("spc", str(o["spc"]))
    return tb


def title(s, t):
    text(s, 0.89, 0.89, 11.9, 0.55, [(t, dict(sz=30, font=SERIF, lh=36))])


def kicker(s, y, t, c=BLUE):
    text(s, 0.89, y, 11.9, 0.28, [(t.upper(), dict(sz=12, b=True, c=c, spc=100, lh=16.8))])


def footer(s, t):
    text(s, 0.89, 6.97, 11.9, 0.3, [(t, dict(sz=11, c=MUTED, lh=15))])


# ---------------------------------------------------------------- Slide 1
s1 = new_slide()
title(s1, "Our framework: from macro themes to a robust portfolio")
kicker(s1, 1.66, "The process · five steps, applied to every theme")

steps = [
    ("01 · Themes", "Identify the themes",
     "12–15 overarching macro themes that drive returns across rates, FX, equities and commodities"),
    ("02 · Probabilities", "Assign distributions",
     "A probability distribution of outcomes for each theme, compared against what markets already price"),
    ("03 · Trades", "Pick the best expression",
     "The trade with the best payout for its upfront cost: premium, carry and relative value"),
    ("04 · Correlations", "Map co-movement",
     "Expected correlations and payoffs of trades across themes: where risk diversifies and where it concentrates"),
    ("05 · Allocation", "Allocate capital",
     "Size by probability-weighted payoff and risk contribution; estimate the PnL distribution of the whole book"),
]
n, x0, W, gap, y, h = 5, 0.89, 11.9, 0.2, 2.01, 2.35
cw = (W - gap * (n - 1)) / n
for i, (lab, head, body) in enumerate(steps):
    x = x0 + i * (cw + gap)
    box(s1, x, y, cw, h)
    text(s1, x + 0.17, y + 0.17, cw - 0.34, h - 0.3, [
        (lab.upper(), dict(sz=11, b=True, c=BLUE, spc=100, lh=15)),
        (head, dict(sz=15, b=True, lh=19, before=6)),
        (body, dict(sz=12, c=GREY, lh=16.5, before=6)),
    ])

# Bottom row: starting themes (left) + the lens (right)
yb, hb = 4.56, 2.05
lw = 6.6
box(s1, x0, yb, lw, hb)
text(s1, x0 + 0.22, yb + 0.18, lw - 0.44, 0.25,
     [("Starting themes · 4 of a targeted 12–15".upper(), dict(sz=11, b=True, c=RED, spc=100, lh=15))])
themes = [
    ("Dollar debasement", "Fiscal and policy risks to the dollar's role as a store of value"),
    ("AI boom", "Capex, productivity and the growth–inflation mix"),
    ("European growth and fiscal", "Defence and infrastructure spending, and Bund supply"),
    ("UK: BoE hikes gradually", "Pace of the tightening cycle against market pricing"),
]
tw = (lw - 0.44 - 0.3) / 2
for i, (head, body) in enumerate(themes):
    tx = x0 + 0.22 + (i % 2) * (tw + 0.3)
    ty = yb + 0.6 + (i // 2) * 0.72
    text(s1, tx, ty, tw, 0.6, [
        (head, dict(sz=13, b=True, lh=17)),
        (body, dict(sz=11, c=GREY, lh=15, before=2)),
    ])

rx = x0 + lw + gap
rw = W - lw - gap
box(s1, rx, yb, rw, hb, fill=NAVY, line=None)
text(s1, rx + 0.25, yb + 0.2, rw - 0.5, hb - 0.4, [
    ("The lens: risk/reward".upper(), dict(sz=11, b=True, c=ONDARK_MUTED, spc=100, lh=15)),
    ("Every theme and trade is judged on what it costs against what it pays, not on conviction alone.",
     dict(sz=14, c=ONDARK, lh=19, before=6)),
    ("Upfront premium · payout/cost ratio across all trades · probability-weighted PnL of the portfolio",
     dict(sz=12, c=ONDARK_MUTED, lh=16.5, before=8)),
])

footer(s1, "Framework in development. Themes shown are an illustrative starting set; no positions or recommendations implied.")
s1.notes_slide.notes_text_frame.text = (
    "The aim is to complement the current strategy with a more systematic portfolio construction process. "
    "We start from a small set of overarching macro themes (12-15), assign each one a probability distribution of outcomes, "
    "and compare it with what markets already price. For each theme we look for the most efficient trade, the best payout for its "
    "upfront cost, then measure how trades co-move across themes before allocating capital. Throughout, the yardstick is risk/reward: "
    "premium paid, payout/cost ratio, and the probability-weighted PnL of the whole book rather than of single trades. "
    "The four themes shown are the starting set; the list will grow to 12-15."
)

# ---------------------------------------------------------------- Slide 2
s2 = new_slide()
title(s2, "How we build it: an analytics toolkit today, AI agents next")
kicker(s2, 1.66, "The toolkit · mapped to the framework")

tools = [
    ("01 · Themes", "Macro regime",
     "Where growth and inflation surprise, and whether rates price it",
     "Macro heatmaps across 10 economies; global yields factor model"),
    ("02 · Probabilities", "Implied odds",
     "What is already priced, so we pay only for views that differ",
     "Option-implied scenario probabilities; macro options monitor; inflation fixings"),
    ("03 · Trades", "Relative value",
     "The cheapest, best-carrying expression of a view",
     "Swap and EGB RV; fitted curves; PCA; mean reversion; seasonality"),
    ("03 · Trades", "Positioning",
     "Crowded trades pay less; positioning helps with timing",
     "Euribor futures positioning; CTA momentum; data-print analysis"),
    ("04–05 · Portfolio", "Correlations",
     "Regime breaks that change how the book diversifies",
     "Correlation regime scanner; trade blotter and trade watch"),
]
y, h = 2.01, 2.45
for i, (lab, head, why, names) in enumerate(tools):
    x = x0 + i * (cw + gap)
    box(s2, x, y, cw, h)
    text(s2, x + 0.17, y + 0.17, cw - 0.34, 1.2, [
        (lab.upper(), dict(sz=11, b=True, c=BLUE, spc=100, lh=15)),
        (head, dict(sz=15, b=True, lh=19, before=6)),
        (why, dict(sz=12, lh=16, before=6)),
    ])
    rule(s2, x + 0.17, y + 1.6, cw - 0.34)
    text(s2, x + 0.17, y + 1.7, cw - 0.34, h - 1.8, [(names, dict(sz=11, c=GREY, lh=14.5))])

kicker(s2, 4.66, "Roadmap · AI agents that run the framework", c=RED)
ya, ha = 5.0, 1.78
road = [
    ("Live", "Fixed Income Pod",
     "Specialist agents for macro, RV, inflation and positioning run the toolkit and draft a joint view"),
    ("Next", "Portfolio Builder agent",
     "Draws on the toolkit, the book and sell-side research to propose probabilities, trades, correlations and sizing"),
    ("Next", "Risk agent",
     "Computes the book's sensitivities to specific risk factors, so any exposure can be hedged quickly"),
]
aw = 2.75
for i, (stage, head, body) in enumerate(road):
    x = x0 + i * (aw + gap)
    box(s2, x, ya, aw, ha)
    text(s2, x + 0.2, ya + 0.17, aw - 0.4, ha - 0.3, [
        [(stage.upper(), dict(sz=11, b=True, c=BLUE if stage == "Live" else RED, spc=100, lh=15))],
        (head, dict(sz=14, b=True, lh=18, before=4)),
        (body, dict(sz=12, c=GREY, lh=16, before=4)),
    ])
dx = x0 + 3 * (aw + gap)
dw = W - (dx - x0)
box(s2, dx, ya, dw, ha, fill=NAVY, line=None)
text(s2, dx + 0.25, ya + 0.2, dw - 0.5, ha - 0.4, [
    ("Why AI".upper(), dict(sz=11, b=True, c=ONDARK_MUTED, spc=100, lh=15)),
    ("More markets and asset classes covered by the same team, with deeper analysis.",
     dict(sz=13, c=ONDARK, lh=17.5, before=6)),
    ("The PM keeps every decision.", dict(sz=12, c=ONDARK_MUTED, lh=16, before=6)),
])

footer(s2, "Tools built in-house; agents are decision support. Portfolio Builder and Risk agents are planned, not yet live.")
s2.notes_slide.notes_text_frame.text = (
    "Each group of tools answers one step of the framework. Macro regime: dynamic factor models across 10 economies show where "
    "data is surprising and whether rates reflect it. Market-implied odds: option-implied distributions tell us what is already in the "
    "price, so we only take risk where our probabilities differ. Relative value and positioning find the cheapest and best-timed "
    "expression of a view. Correlation regimes and the trade blotter monitor how the book diversifies. "
    "On the roadmap, the Fixed Income Pod already runs these tools through specialist AI agents. Next, a Portfolio Builder agent will "
    "carry out the full framework using the toolkit, the live portfolio and sell-side research, and a Risk agent will compute the book's "
    "Greeks to specific risk factors so we can hedge quickly. The PM retains all decisions."
)

prs.save(OUT)
print("saved", OUT)
