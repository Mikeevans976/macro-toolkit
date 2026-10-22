"""US GDP nowcast (quarterly target in the Quant Macro engine), GDPNow and FOMC SEP comparisons.

The model is fitted once on the simulated US panel with few EM iterations; the checks are
identities that hold for any parameter values. Public downloads are switched off and saved
state goes to a temporary directory.
"""
from __future__ import annotations

import dataclasses
import io
import warnings
from datetime import datetime

import numpy as np
import openpyxl
import pandas as pd
import pytest
import time_machine
from fastapi.testclient import TestClient

import fed_data
import nowcast as nc
import quant_macro as qm

AS_OF = pd.Timestamp("2026-10-22")
SPEC = dataclasses.replace(qm.MODELS["us-gdp"], maxiter=20, backtest_from="2025Q1", horizons=(-30, -1))


@pytest.fixture(scope="module", autouse=True)
def state_dir(tmp_path_factory):
    saved = qm.ROOT, qm.FETCH_PUBLIC
    qm.ROOT = tmp_path_factory.mktemp("quant_macro_us")
    qm.FETCH_PUBLIC = False
    yield qm.ROOT
    qm.ROOT, qm.FETCH_PUBLIC = saved


@pytest.fixture(scope="module")
def fitted() -> qm.Fitted:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return qm.fit(SPEC, AS_OF, "csv")


@pytest.fixture(scope="module")
def payload(fitted) -> dict:
    return qm.build_payload(fitted)


def test_simulated_gdp_is_quarterly_and_the_triangle_of_its_monthly_path():
    panel = nc.simulate_panel(SPEC, AS_OF)
    g = panel.values["gdp"].dropna()
    assert all(p.month % 3 == 0 for p in g.index)
    assert g.index.max() <= nc.horizon_end(AS_OF, SPEC.forecast_quarters)
    # as of the day, the last published quarter is 2026Q2 (Q3's advance estimate is due 29 Oct)
    assert panel.vintage(AS_OF)["gdp"].last_valid_index() == pd.Period("2026-06", "M")


def test_quarterly_series_go_last_and_reapplying_keeps_the_structure(fitted):
    model = fitted.res.model
    assert list(model.endog_names)[-1] == "gdp" and model.k_endog_M == model.k_endog - 1
    again = fitted.at(AS_OF)
    assert again.model.k_states == model.k_states and again.model.k_endog_M == model.k_endog_M
    q = pd.Period("2026Q3", "Q")
    assert nc.quarter_nowcast(again, "gdp", q, "Q") == pytest.approx(nc.quarter_nowcast(fitted.res, "gdp", q, "Q"))
    # the quarterly design row carries the triangle: its lags span five months
    z = nc._matrix(fitted.res, "design")[model.k_endog - 1]
    assert np.count_nonzero(np.abs(z) > 1e-12) > 5


def test_payload_quarters_and_next_print(payload):
    rows = {r["quarter"]: r for r in payload["quarters"]}
    assert payload["target_freq"] == "Q" and payload["target_units"] == "% q/q saar"
    assert rows["2026Q2"]["status"] == "released" and rows["2026Q2"]["actual"] is not None
    assert rows["2026Q3"]["status"] == "backcast" and rows["2026Q4"]["status"] == "nowcast"
    assert rows["2027Q4"]["status"] == "forecast"                      # four quarters past the current one
    assert all(r["months"] == [] for r in payload["quarters"]) and payload["monthly_path"] == []
    n = payload["next_print"]
    assert n["period"] == "2026-09" and n["quarter"] is None and n["model"] == rows["2026Q3"]["nowcast"]
    assert "3Q 2026" in n["label"]
    # news is tracked for the quarters up to next quarter only (the current one is 2026Q4)
    assert set(payload["evolution"][-1]["nowcasts"]) == {"2026Q3", "2026Q4", "2027Q1"}


def test_news_adds_up_for_a_quarterly_target(payload):
    steps = payload["evolution"]
    for prev, step in zip(steps, steps[1:]):
        for q in step["nowcasts"]:
            total = sum((r["impacts"][q] or 0.0) for r in step["releases"])
            assert step["nowcasts"][q] - prev["nowcasts"][q] == pytest.approx(total, abs=2e-3)


def test_policy_comparison_compounds_quarters(payload):
    pol = payload["policy"]
    assert pol["label"] == "FOMC SEP median" and pol["gdp"]["2026"] == 2.3 and pol["unemployment"]["2026"] == 4.1
    rows = {r["quarter"]: (r["actual"] if r["actual"] is not None else r["nowcast"]) for r in payload["quarters"]}
    v = [rows[f"2027Q{i}"] for i in range(1, 5)]
    expected = 100 * (np.prod([(1 + x / 100) ** 0.25 for x in v]) - 1)
    y2027 = next(y for y in pol["model_years"] if y["year"] == 2027)
    assert y2027["model"] == pytest.approx(expected, abs=1e-3) and y2027["quarters_published"] == 0
    assert next(y for y in pol["model_years"] if y["year"] == 2028)["model"] is None
    assert payload["central_bank"] is None                              # downloads off in tests


def test_parse_gdpnow(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("TrackingArchives")
    ws.append(["Forecast Date", "Quarter being forecasted"] + [f"c{i}" for i in range(25)] + ["GDP Nowcast"])
    ws.append([datetime(2026, 5, 1), datetime(2026, 6, 30)] + [0.0] * 25 + [1.8])
    ws.append([datetime(2026, 7, 28), datetime(2026, 6, 30)] + [0.0] * 25 + [1.54])
    ws = wb.create_sheet("TrackingHistory")
    ws.append([None, None, datetime(2026, 7, 30), datetime(2026, 10, 20)])
    ws.append(["Evolution of GDP nowcast and components for 2026q3"])
    ws.append(["GDP", "GDP Nowcast", 4.95, 3.5])
    buf = io.BytesIO()
    wb.save(buf)
    (tmp_path / "g.xlsx").write_bytes(buf.getvalue())
    est = fed_data.parse_gdpnow(tmp_path / "g.xlsx", today=pd.Timestamp("2026-10-22"))
    assert [(e["published"], e["quarter"], e["value"]) for e in est] == [
        ("2026-05-01", "2026Q2", 1.8), ("2026-07-28", "2026Q2", 1.54), ("2026-07-30", "2026Q3", 4.95), ("2026-10-20", "2026Q3", 3.5)]


def test_backtest_scores_a_quarterly_target():
    spec = dataclasses.replace(SPEC, backtest_from="2026Q2", horizons=(-1,))       # one quarter, one origin
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        bt = qm.build_backtest(spec, nc.simulate_panel(spec, AS_OF), AS_OF)
    assert [r["quarter"] for r in bt["rows"]] == ["2026Q2"]
    r = bt["rows"][0]
    assert r["actual"] is not None and r["dfm"][0] is not None and r["release_date"] == "2026-07-29"


@pytest.fixture(scope="module")
def client(fitted, payload):
    import main
    from auth import get_current_user
    main.app.dependency_overrides[get_current_user] = lambda: {"username": "t", "full_name": "T"}
    saved = dict(qm.MODELS)
    qm.MODELS["us-gdp"] = SPEC
    # the module's fit, saved: the route serves it without estimating again
    qm._write(qm._dir("us-gdp") / "model.json", fitted.state)
    qm._write(qm._dir("us-gdp") / "nowcast.json", payload)
    with warnings.catch_warnings(), time_machine.travel("2026-10-22T12:00:00+00:00", tick=False):
        warnings.simplefilter("ignore")
        yield TestClient(main.app)
    qm.MODELS.clear()
    qm.MODELS.update(saved)
    main.app.dependency_overrides.clear()


def test_api_lists_and_serves_the_us_model(client, fitted):
    slugs = [m["slug"] for m in client.get("/api/tools/quant-macro").json()]
    assert slugs == ["uk-gdp", "us-gdp"]
    r = client.get("/api/tools/quant-macro/us-gdp")
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["target_freq"] == "Q" and body["policy"]["release"] == "September 2026"
    assert client.get("/api/tools/quant-macro/us-gdp/status").json()["model_fitted_at"] == fitted.state["fitted_at"]
