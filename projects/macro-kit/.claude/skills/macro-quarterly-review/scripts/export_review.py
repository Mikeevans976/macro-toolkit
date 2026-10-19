"""Export the quarterly review markdown to .docx (python-docx) and .html (for Edge PDF printing)."""
import re
import sys

import os
if os.environ.get("REVIEW_LIBS"):
    sys.path.insert(0, os.environ["REVIEW_LIBS"])  # folder with python-docx and markdown installed via pip --target
import markdown
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC, DOCX, HTML = sys.argv[1], sys.argv[2], sys.argv[3]
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x5F, 0x6D)
lines = open(SRC, encoding="utf-8").read().splitlines()


# ---------- DOCX ----------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(table, cm=0.08):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(cm * 567))); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def add_runs(par, text, size=None, color=None, bold=False):
    # split on **bold** and *italic*
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        b, i = bold, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            tok, i = tok[1:-1], True
        r = par.add_run(tok)
        r.bold, r.italic = b, i
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.orientation = WD_ORIENT.PORTRAIT
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(1.7))
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)

st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.1
for name, size, color in (("Title", 22, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 11.5, NAVY)):
    s = doc.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.color.rgb = color; s.font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Quarterly Review: Q3 2026  ·  page "); r.font.size = Pt(8); r.font.color.rgb = GREY
for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run(); run.font.size = Pt(8); run.font.color.rgb = GREY
    if kind:
        f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); run._r.append(f)
    else:
        t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = txt; run._r.append(t)

i, first_h2 = 0, True
while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "):
        p = doc.add_paragraph(style="Title"); add_runs(p, ln[2:], color=NAVY)
    elif ln.startswith("## "):
        if not first_h2:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first_h2 = False
        doc.add_heading(ln[3:], level=1)
    elif ln.startswith("### "):
        doc.add_heading(ln[4:], level=2)
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-\s|:]+\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        ncol = len(rows[0])
        t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(t)
        wide = ncol >= 8
        for ri, row in enumerate(rows):
            total = any(c.startswith("**") for c in row[:2])
            for ci in range(ncol):
                c = t.cell(ri, ci); c.text = ""
                par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
                txt = row[ci] if ci < len(row) else ""
                if ri == 0:
                    add_runs(par, txt, size=7.5 if wide else 8.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True); shade(c, "1F3A5F")
                else:
                    add_runs(par, txt, size=7.5 if wide else 8.5, bold=total)
                    if total: shade(c, "DCE4EE")
                    elif ri % 2 == 0: shade(c, "F3F5F8")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    elif re.match(r"^\s*- ", ln):
        level = (len(ln) - len(ln.lstrip(" "))) // 2
        p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, ln.strip()[2:])
    elif ln.strip() == "---" or not ln.strip():
        pass
    else:
        p = doc.add_paragraph()
        if ln.startswith("*") and not ln.startswith("**"):
            add_runs(p, ln, size=9, color=GREY)
        else:
            add_runs(p, ln)
    i += 1
doc.save(DOCX)

# ---------- HTML (for PDF) ----------
md = []
for l in lines:
    if l.strip() == "---":
        continue
    if re.match(r"^- ", l) and md and md[-1].strip() and not re.match(r"^\s*- ", md[-1]):
        md.append("")  # python-markdown needs a blank line before a list
    md.append(re.sub(r"^  - ", "    - ", l))  # nested bullets need 4-space indent
body = markdown.markdown("\n".join(md), extensions=["tables", "sane_lists"])
# bold the total rows
def mark(m):
    row = m.group(0)
    cells = re.findall(r"<td>(.*?)</td>", row, flags=re.S)
    return row.replace("<tr>", '<tr class="total">', 1) if any("<strong>" in c for c in cells[:2]) else row
body = re.sub(r"<tr>.*?</tr>", mark, body, flags=re.S)
body = re.sub(r"<table>(?=\s*<thead>\s*<tr>(?:(?!</tr>).)*Exp\. P&amp;L)", '<table class="trades">', body, flags=re.S)
css = """
@page { size: A4; margin: 15mm 14mm 16mm 14mm; }
body { font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 9.6pt; line-height: 1.38; color: #1d2430; }
h1 { color: #1f3a5f; font-size: 22pt; margin: 0 0 4pt; }
h1 + p { color: #555f6d; font-size: 8.8pt; border-bottom: 2px solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 10pt; }
h2 { color: #1f3a5f; font-size: 14.5pt; margin: 0 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid #c9d3df; }
h2:not(:first-of-type) { break-before: page; }
h3 { color: #1f3a5f; font-size: 11pt; margin: 12pt 0 4pt; }
p { margin: 4pt 0; }
ul { margin: 3pt 0 6pt; padding-left: 16pt; } li { margin: 1.5pt 0; } ul ul { margin: 1pt 0; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 8pt; font-size: 7.9pt; break-inside: avoid; }
th { background: #1f3a5f; color: #fff; text-align: left; padding: 3pt 4pt; font-weight: 600; }
td { padding: 2.6pt 4pt; border-bottom: 1px solid #dfe5ec; vertical-align: top; }
tr:nth-child(even) td { background: #f3f5f8; }
tr.total td { background: #dce4ee; font-weight: 600; }
table.trades td:last-child, table.trades td:nth-last-child(2) { white-space: nowrap; }
"""
open(HTML, "w", encoding="utf-8").write(f"<!doctype html><html><head><meta charset='utf-8'><title>Quarterly Review: Q3 2026</title><style>{css}</style></head><body>{body}</body></html>")
print("ok")
